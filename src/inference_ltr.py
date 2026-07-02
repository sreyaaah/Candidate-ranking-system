import numpy as np
import pandas as pd
import sqlite3
import os
import sys
import xgboost as xgb

# Ensure src is in the path to import search_hybrid
sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from search_hybrid import get_hybrid_scores
from docx import Document

required_skills = [
    "python",
    "react",
    "node.js",
    "docker",
    "sql"
]

def main():
    print("Loading Job Description...")
    doc = Document("data/jobs/job_description.docx")
    jd_text = "\n".join([para.text for para in doc.paragraphs])
    
    # 1. Stage 1 & 2: Hybrid Retrieval
    top_n = 2000
    hybrid_results = get_hybrid_scores(jd_text, top_n=top_n)
    
    print("\nFetching features from database...")
    candidate_ids = [res[0] for res in hybrid_results]
    
    # 2. Stage 3: Feature Engineering lookups
    conn = sqlite3.connect("data/candidates.db")
    cursor = conn.cursor()
    
    placeholders = ",".join(["?"] * len(candidate_ids))
    query = f"""
        SELECT 
            c.id, c.skills, 
            f.yoe, f.ai_years, f.job_hopping_index, f.github_score, 
            f.education_tier, f.notice_period_days
        FROM candidates c
        LEFT JOIN candidate_features f ON c.id = f.candidate_id
        WHERE c.id IN ({placeholders})
    """
    cursor.execute(query, candidate_ids)
    
    features_map = {}
    for row in cursor.fetchall():
        features_map[row[0]] = {
            "skills": row[1].lower() if row[1] else "",
            "yoe": row[2] or 0.0,
            "ai_years": row[3] or 0.0,
            "job_hopping_index": row[4] or 0.0,
            "github_score": row[5] or 0.0,
            "education_tier": row[6] or 3,
            "notice_period": row[7] or 0
        }
        
    conn.close()
    
    # Prepare DataFrame for XGBoost Inference
    dataset = []
    
    for cand_id, rrf_score in hybrid_results:
        feats = features_map.get(cand_id, {})
        text = feats.get("skills", "")
        
        matched = sum(1 for skill in required_skills if skill in text)
        skill_score = matched / len(required_skills)
            
        dataset.append({
            "candidate_id": cand_id,
            "rrf_score": rrf_score,
            "skill_score": skill_score,
            "yoe": feats.get("yoe", 0.0),
            "ai_years": feats.get("ai_years", 0.0),
            "job_hopping_index": feats.get("job_hopping_index", 0.0),
            "github_score": feats.get("github_score", 0.0),
            "education_tier": feats.get("education_tier", 3),
            "notice_period": feats.get("notice_period", 0)
        })
        
    df = pd.DataFrame(dataset)
    
    # 3. Stage 4: Learning-to-Rank Inference
    print("Loading XGBoost Ranker model...")
    ranker = xgb.XGBRanker()
    ranker.load_model("models/xgb_ranker.json")
    
    features = [
        "rrf_score", "skill_score", "yoe", "ai_years", 
        "job_hopping_index", "github_score", "education_tier", "notice_period"
    ]
    
    print("Re-scoring candidates using ML model...")
    X = df[features]
    predictions = ranker.predict(X)
    
    df["ml_score"] = predictions
    
    # Sort by ML score
    df = df.sort_values(by="ml_score", ascending=False).reset_index(drop=True)
    
    print("\nFinal ML Candidate Ranking (Top 20):\n")
    for i in range(min(20, len(df))):
        c = df.iloc[i]
        print(f"{i+1}. {c['candidate_id']}")
        print(f"   ML Rank Score: {c['ml_score']:.4f}")
        print(f"   YoE: {c['yoe']:.1f} | Edu: Tier {c['education_tier']} | Notice: {c['notice_period']}d | RRF: {c['rrf_score']:.4f}\n")

if __name__ == "__main__":
    main()
