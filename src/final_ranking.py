import numpy as np
import sqlite3
import os
import sys

# Ensure src is in the path to import search_hybrid
sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from search_hybrid import get_hybrid_scores
from docx import Document

required_skills = [
    "python",
    "react",
    "node.js",
    "docker",
    "sql"
]

def main():
    print("Loading Job Description...")
    doc = Document("data/jobs/job_description.docx")
    jd_text = "\n".join([para.text for para in doc.paragraphs])
    
    # Get Hybrid RRF scores
    top_n = 2000
    hybrid_results = get_hybrid_scores(jd_text, top_n=top_n)
    
    print("\nFetching skills and features from database...")
    candidate_ids = [res[0] for res in hybrid_results]
    
    conn = sqlite3.connect("data/candidates.db")
    cursor = conn.cursor()
    
    placeholders = ",".join(["?"] * len(candidate_ids))
    query = f"""
        SELECT 
            c.id, c.skills, 
            f.yoe, f.github_score, f.education_tier, f.notice_period_days, f.ai_years
        FROM candidates c
        LEFT JOIN candidate_features f ON c.id = f.candidate_id
        WHERE c.id IN ({placeholders})
    """
    cursor.execute(query, candidate_ids)
    
    features_map = {}
    for row in cursor.fetchall():
        features_map[row[0]] = {
            "skills_text": row[1].lower() if row[1] else "",
            "yoe": row[2] or 0.0,
            "github_score": row[3] or 0.0,
            "education_tier": row[4] or 3,
            "notice_period": row[5] or 0,
            "ai_years": row[6] or 0.0
        }
        
    conn.close()
    
    final_scores = []
    
    for cand_id, rrf_score in hybrid_results:
        feats = features_map.get(cand_id, {})
        text = feats.get("skills_text", "")
        
        # Skill calculation
        matched = sum(1 for skill in required_skills if skill in text)
        skill_score = matched / len(required_skills)
        
        # Scaled RRF Score
        scaled_rrf = rrf_score * 30
        
        # Base Final Score
        final_score = (0.7 * scaled_rrf) + (0.3 * skill_score)
        
        # --- Stage 3 Heuristic Feature Boosts & Penalties ---
        yoe = feats.get("yoe", 0.0)
        github = feats.get("github_score", 0.0)
        tier = feats.get("education_tier", 3)
        notice = feats.get("notice_period", 0)
        
        # Small boost for high YoE (capped at 10 years for boost)
        final_score += min(yoe, 10) * 0.01 
        
        # Small boost for GitHub activity
        final_score += (github / 100.0) * 0.1
        
        # Boost for Tier 1 institutions
        if tier == 1:
            final_score += 0.05
            
        # Heavy penalty for very long notice periods (e.g., > 60 days)
        if notice > 60:
            final_score -= 0.10
            
        final_scores.append({
            "id": cand_id,
            "rrf": rrf_score,
            "scaled_rrf": scaled_rrf,
            "skill": skill_score,
            "final": final_score,
            "yoe": yoe,
            "github": github,
            "tier": tier,
            "notice": notice
        })
        
    # Re-sort by the final combined score
    final_scores.sort(key=lambda x: x["final"], reverse=True)
    
    print("\nFinal Candidate Ranking (Top 20 with Features):\n")
    for i in range(min(20, len(final_scores))):
        c = final_scores[i]
        print(f"{i+1}. {c['id']}")
        print(f"   RRF (Scaled): {c['scaled_rrf']:.4f} | Skill Score: {c['skill']:.2f}")
        print(f"   YoE: {c['yoe']:.1f} | GitHub: {c['github']:.1f} | Edu Tier: {c['tier']} | Notice: {c['notice']} days")
        print(f"   FINAL HEURISTIC SCORE: {c['final']:.4f}\n")

if __name__ == "__main__":
    main()