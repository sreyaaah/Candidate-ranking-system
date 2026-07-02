import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    
    # Set slide dimensions to widescreen (16:9)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Color Palette Constants
    DARK_BG = RGBColor(11, 15, 25)
    WHITE = RGBColor(243, 244, 246)
    LIGHT_GRAY = RGBColor(156, 163, 175)
    PURPLE = RGBColor(167, 139, 250)
    BLUE = RGBColor(59, 130, 246)
    CARD_BG = RGBColor(20, 26, 40)
    BORDER_COLOR = RGBColor(55, 65, 81)
    
    blank_layout = prs.slide_layouts[6] # Blank layout
    
    # ----------------------------------------------------
    # SLIDE 1: Title & System Overview
    # ----------------------------------------------------
    slide_1 = prs.slides.add_slide(blank_layout)
    
    # Add dark background
    bg_shape_1 = slide_1.shapes.add_shape(
        1, Inches(0), Inches(0), Inches(13.333), Inches(7.5) # 1 = rectangle
    )
    bg_shape_1.fill.solid()
    bg_shape_1.fill.fore_color.rgb = DARK_BG
    bg_shape_1.line.fill.background() # No border
    
    # Title & Subtitle box
    title_box_1 = slide_1.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.5))
    tf_1 = title_box_1.text_frame
    tf_1.word_wrap = True
    
    p_badge = tf_1.paragraphs[0]
    p_badge.text = "SLIDE 1  |  SYSTEM OVERVIEW"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = PURPLE
    p_badge.font.name = "Arial"
    p_badge.space_after = Pt(10)
    
    p_title = tf_1.add_paragraph()
    p_title.text = "Modern Multi-Stage Hybrid Candidate Ranking Funnel"
    p_title.font.size = Pt(36)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    p_title.font.name = "Arial"
    
    # Left Column: The Architecture Concept
    left_col_1 = slide_1.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(5.3), Inches(4.0))
    tf_left_1 = left_col_1.text_frame
    tf_left_1.word_wrap = True
    
    p_concept_title = tf_left_1.paragraphs[0]
    p_concept_title.text = "Solving the Recruitment Search Paradox"
    p_concept_title.font.size = Pt(20)
    p_concept_title.font.bold = True
    p_concept_title.font.color.rgb = PURPLE
    p_concept_title.font.name = "Arial"
    p_concept_title.space_after = Pt(12)
    
    p_concept_desc = tf_left_1.add_paragraph()
    p_concept_desc.text = "Traditional recruitment software relies on simple exact keyword matching (which misses context and synonyms) or dense semantic search (which dilutes niche technical terms). This project builds an advanced multi-stage ranking pipeline that combines both styles with machine learning LTR features:"
    p_concept_desc.font.size = Pt(14)
    p_concept_desc.font.color.rgb = LIGHT_GRAY
    p_concept_desc.font.name = "Arial"
    p_concept_desc.space_after = Pt(14)
    
    bullets = [
        "Fuses FAISS dense semantic indexes with BM25 sparse lexical matching.",
        "Applies late-fusion Reciprocal Rank Fusion (RRF) to merge candidate pools.",
        "Employs dual-objective XGBoost Learning-to-Rank student models.",
        "Leverages neural Cross-Encoder matching for top candidates.",
        "Enforces Maximal Marginal Relevance (MMR) for result diversity."
    ]
    for bullet in bullets:
        p_b = tf_left_1.add_paragraph()
        p_b.text = "• " + bullet
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = WHITE
        p_b.font.name = "Arial"
        p_b.space_after = Pt(6)
        
    # Right Column: Funnel Box Visual Representation
    right_col_1 = slide_1.shapes.add_textbox(Inches(7.0), Inches(2.6), Inches(5.3), Inches(4.0))
    tf_right_1 = right_col_1.text_frame
    tf_right_1.word_wrap = True
    
    p_funnel_title = tf_right_1.paragraphs[0]
    p_funnel_title.text = "Pipeline Selection Funnel"
    p_funnel_title.font.size = Pt(20)
    p_funnel_title.font.bold = True
    p_funnel_title.font.color.rgb = BLUE
    p_funnel_title.font.name = "Arial"
    p_funnel_title.space_after = Pt(20)
    
    stages = [
        ("1. Raw Candidates (candidates.jsonl)", "N = 100,000"),
        ("2. RRF Late-Fusion Retrieval Pool", "N = 2,000"),
        ("3. Blended XGBoost LTR Inference", "N = 200"),
        ("4. Cross-Encoder + MMR Final Output", "N = 100")
    ]
    
    for label, count in stages:
        p_stage = tf_right_1.add_paragraph()
        p_stage.text = f"{label}  --->  {count}"
        p_stage.font.size = Pt(14)
        p_stage.font.bold = True
        p_stage.font.color.rgb = WHITE
        p_stage.font.name = "Arial"
        p_stage.space_after = Pt(16)
        
    # ----------------------------------------------------
    # SLIDE 2: Data Preprocessing & Offline Distillation
    # ----------------------------------------------------
    slide_2 = prs.slides.add_slide(blank_layout)
    
    # Add dark background
    bg_shape_2 = slide_2.shapes.add_shape(
        1, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg_shape_2.fill.solid()
    bg_shape_2.fill.fore_color.rgb = DARK_BG
    bg_shape_2.line.fill.background()
    
    # Title & Subtitle box
    title_box_2 = slide_2.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.5))
    tf_2 = title_box_2.text_frame
    tf_2.word_wrap = True
    
    p_badge_2 = tf_2.paragraphs[0]
    p_badge_2.text = "SLIDE 2  |  DATA PREPROCESSING & DISTILLATION"
    p_badge_2.font.size = Pt(12)
    p_badge_2.font.bold = True
    p_badge_2.font.color.rgb = PURPLE
    p_badge_2.font.name = "Arial"
    p_badge_2.space_after = Pt(10)
    
    p_title_2 = tf_2.add_paragraph()
    p_title_2.text = "Offline Feature Engineering & Knowledge Distillation"
    p_title_2.font.size = Pt(36)
    p_title_2.font.bold = True
    p_title_2.font.color.rgb = WHITE
    p_title_2.font.name = "Arial"
    
    # Left Column: Tabular Feature Engineering
    left_col_2 = slide_2.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(5.3), Inches(4.0))
    tf_left_2 = left_col_2.text_frame
    tf_left_2.word_wrap = True
    
    p_eng_title = tf_left_2.paragraphs[0]
    p_eng_title.text = "1. Tabular Feature Engineering"
    p_eng_title.font.size = Pt(20)
    p_eng_title.font.bold = True
    p_eng_title.font.color.rgb = PURPLE
    p_eng_title.font.name = "Arial"
    p_eng_title.space_after = Pt(12)
    
    p_eng_desc = tf_left_2.add_paragraph()
    p_eng_desc.text = "Raw profiles are extracted from JSONL into a clean SQLite schema. We compute and engineer 46 advanced tabular features for every profile to feed structural context into XGBoost:"
    p_eng_desc.font.size = Pt(14)
    p_eng_desc.font.color.rgb = LIGHT_GRAY
    p_eng_desc.font.name = "Arial"
    p_eng_desc.space_after = Pt(12)
    
    eng_bullets = [
        "Career continuity ratios & promotion velocity indicators.",
        "Stability scores (job counts vs. total years of experience).",
        "Education tier checks (extracting Tier 1 institutions).",
        "Logistical flags (salary expectations and notice periods)."
    ]
    for bullet in eng_bullets:
        p_b = tf_left_2.add_paragraph()
        p_b.text = "• " + bullet
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = WHITE
        p_b.font.name = "Arial"
        p_b.space_after = Pt(6)
        
    # Right Column: Knowledge Distillation & Defense
    right_col_2 = slide_2.shapes.add_textbox(Inches(7.0), Inches(2.6), Inches(5.3), Inches(4.0))
    tf_right_2 = right_col_2.text_frame
    tf_right_2.word_wrap = True
    
    p_dist_title = tf_right_2.paragraphs[0]
    p_dist_title.text = "2. Distillation & Adversarial Defense"
    p_dist_title.font.size = Pt(20)
    p_dist_title.font.bold = True
    p_dist_title.font.color.rgb = BLUE
    p_dist_title.font.name = "Arial"
    p_dist_title.space_after = Pt(12)
    
    p_dist_desc = tf_right_2.add_paragraph()
    p_dist_desc.text = "Candidate databases lack explicit training labels. We use knowledge distillation from a neural model and enforce adversarial security:"
    p_dist_desc.font.size = Pt(14)
    p_dist_desc.font.color.rgb = LIGHT_GRAY
    p_dist_desc.font.name = "Arial"
    p_dist_desc.space_after = Pt(12)
    
    dist_bullets = [
        "Neural Teacher: High-capacity Cross-Encoder scores 1500 profiles across 4 query groups to label training data.",
        "Quantile Binning: Continuous scores are binned into classes [0, 1, 2, 3] representing mismatch to outstanding alignment.",
        "Honeypot Hardening: Profiles containing hidden prompt injections or keyword-stuffing flags are flagged.",
        "Adversarial Penalty: Detected honeypots are forced to 0 relevance, training the model to penalize exploits."
    ]
    for bullet in dist_bullets:
        p_b = tf_right_2.add_paragraph()
        p_b.text = "• " + bullet
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = WHITE
        p_b.font.name = "Arial"
        p_b.space_after = Pt(6)
        
    # ----------------------------------------------------
    # SLIDE 3: Online Inference & Decision Engine
    # ----------------------------------------------------
    slide_3 = prs.slides.add_slide(blank_layout)
    
    # Add dark background
    bg_shape_3 = slide_3.shapes.add_shape(
        1, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg_shape_3.fill.solid()
    bg_shape_3.fill.fore_color.rgb = DARK_BG
    bg_shape_3.line.fill.background()
    
    # Title & Subtitle box
    title_box_3 = slide_3.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.5))
    tf_3 = title_box_3.text_frame
    tf_3.word_wrap = True
    
    p_badge_3 = tf_3.paragraphs[0]
    p_badge_3.text = "SLIDE 3  |  ONLINE INFERENCE & DECISION ENGINE"
    p_badge_3.font.size = Pt(12)
    p_badge_3.font.bold = True
    p_badge_3.font.color.rgb = PURPLE
    p_badge_3.font.name = "Arial"
    p_badge_3.space_after = Pt(10)
    
    p_title_3 = tf_3.add_paragraph()
    p_title_3.text = "The Real-Time Ranking & Explainability Engine"
    p_title_3.font.size = Pt(36)
    p_title_3.font.bold = True
    p_title_3.font.color.rgb = WHITE
    p_title_3.font.name = "Arial"
    
    # Left Column: Dual-Objective & Blending
    left_col_3 = slide_3.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(5.3), Inches(4.0))
    tf_left_3 = left_col_3.text_frame
    tf_left_3.word_wrap = True
    
    p_blend_title = tf_left_3.paragraphs[0]
    p_blend_title.text = "1. Ensemble Blending & Scoring"
    p_blend_title.font.size = Pt(20)
    p_blend_title.font.bold = True
    p_blend_title.font.color.rgb = PURPLE
    p_blend_title.font.name = "Arial"
    p_blend_title.space_after = Pt(12)
    
    p_blend_desc = tf_left_3.add_paragraph()
    p_blend_desc.text = "The online engine fetches candidate features and passes them to our blended model ensemble:"
    p_blend_desc.font.size = Pt(14)
    p_blend_desc.font.color.rgb = LIGHT_GRAY
    p_blend_desc.font.name = "Arial"
    p_blend_desc.space_after = Pt(12)
    
    blend_bullets = [
        "XGBoost LTR Ensemble: Combines a Pairwise Ranker (minimizes relative swaps) and NDCG Ranker (optimizes top ranks).",
        "Semantic Fusion: The top 200 LTR candidates are passed to the Cross-Encoder model on CPU.",
        "Final Scoring: Combines LTR (structural features) and Cross-Encoder (text semantic relevance) using a 50/50 blend."
    ]
    for bullet in blend_bullets:
        p_b = tf_left_3.add_paragraph()
        p_b.text = "• " + bullet
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = WHITE
        p_b.font.name = "Arial"
        p_b.space_after = Pt(6)
        
    # Right Column: Diversity & Explainability
    right_col_3 = slide_3.shapes.add_textbox(Inches(7.0), Inches(2.6), Inches(5.3), Inches(4.0))
    tf_right_3 = right_col_3.text_frame
    tf_right_3.word_wrap = True
    
    p_div_title = tf_right_3.paragraphs[0]
    p_div_title.text = "2. Diversity Filter & Explainability"
    p_div_title.font.size = Pt(20)
    p_div_title.font.bold = True
    p_div_title.font.color.rgb = BLUE
    p_div_title.font.name = "Arial"
    p_div_title.space_after = Pt(12)
    
    p_div_desc = tf_right_3.add_paragraph()
    p_div_desc.text = "Post-processing formats the top 100 candidate slate to ensure it is diverse, non-redundant, and explainable:"
    p_div_desc.font.size = Pt(14)
    p_div_desc.font.color.rgb = LIGHT_GRAY
    p_div_desc.font.name = "Arial"
    p_div_desc.space_after = Pt(12)
    
    div_bullets = [
        "MMR Diversity Filter (lambda=0.85): Penalizes candidate profiles that are highly similar/redundant to already selected ones.",
        "Tie-Breaker Sorting: Strictly sorts the slate by score DESC and candidate ID ASC per the hackathon specification.",
        "Structured Explanations: Formulates explanations listing matched/missing skills, highlights, risks, and confidence."
    ]
    for bullet in div_bullets:
        p_b = tf_right_3.add_paragraph()
        p_b.text = "• " + bullet
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = WHITE
        p_b.font.name = "Arial"
        p_b.space_after = Pt(6)
        
    # ----------------------------------------------------
    # SLIDE 4: Results & Performance
    # ----------------------------------------------------
    slide_4 = prs.slides.add_slide(blank_layout)
    
    # Add dark background
    bg_shape_4 = slide_4.shapes.add_shape(
        1, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg_shape_4.fill.solid()
    bg_shape_4.fill.fore_color.rgb = DARK_BG
    bg_shape_4.line.fill.background()
    
    # Title & Subtitle box
    title_box_4 = slide_4.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.5))
    tf_4 = title_box_4.text_frame
    tf_4.word_wrap = True
    
    p_badge_4 = tf_4.paragraphs[0]
    p_badge_4.text = "SLIDE 4  |  EVALUATION METRICS & RESULTS"
    p_badge_4.font.size = Pt(12)
    p_badge_4.font.bold = True
    p_badge_4.font.color.rgb = PURPLE
    p_badge_4.font.name = "Arial"
    p_badge_4.space_after = Pt(10)
    
    p_title_4 = tf_4.add_paragraph()
    p_title_4.text = "Results & Performance"
    p_title_4.font.size = Pt(36)
    p_title_4.font.bold = True
    p_title_4.font.color.rgb = WHITE
    p_title_4.font.name = "Arial"
    
    # Left Column: Results
    left_col_4 = slide_4.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(5.3), Inches(4.0))
    tf_left_4 = left_col_4.text_frame
    tf_left_4.word_wrap = True
    
    p_res_title = tf_left_4.paragraphs[0]
    p_res_title.text = "Results"
    p_res_title.font.size = Pt(20)
    p_res_title.font.bold = True
    p_res_title.font.color.rgb = PURPLE
    p_res_title.font.name = "Arial"
    p_res_title.space_after = Pt(16)
    
    res_bullets = [
        "Successfully ranks candidates according to JD relevance.",
        "Hybrid ranking improves prioritization.",
        "Transparent scoring enables recruiter validation."
    ]
    for bullet in res_bullets:
        p_b = tf_left_4.add_paragraph()
        p_b.text = "• " + bullet
        p_b.font.size = Pt(14)
        p_b.font.color.rgb = WHITE
        p_b.font.name = "Arial"
        p_b.space_after = Pt(12)
        
    # Right Column: Performance
    right_col_4 = slide_4.shapes.add_textbox(Inches(7.0), Inches(2.6), Inches(5.3), Inches(4.0))
    tf_right_4 = right_col_4.text_frame
    tf_right_4.word_wrap = True
    
    p_perf_title = tf_right_4.paragraphs[0]
    p_perf_title.text = "Performance"
    p_perf_title.font.size = Pt(20)
    p_perf_title.font.bold = True
    p_perf_title.font.color.rgb = BLUE
    p_perf_title.font.name = "Arial"
    p_perf_title.space_after = Pt(16)
    
    perf_bullets = [
        "FAISS enables fast similarity search.",
        "Embeddings generated once and reused.",
        "Lightweight Streamlit interface for quick interaction."
    ]
    for bullet in perf_bullets:
        p_b = tf_right_4.add_paragraph()
        p_b.text = "• " + bullet
        p_b.font.size = Pt(14)
        p_b.font.color.rgb = WHITE
        p_b.font.name = "Arial"
        p_b.space_after = Pt(12)

    # ----------------------------------------------------
    # SLIDE 5: Technologies Used
    # ----------------------------------------------------
    slide_5 = prs.slides.add_slide(blank_layout)
    
    # Add dark background
    bg_shape_5 = slide_5.shapes.add_shape(
        1, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg_shape_5.fill.solid()
    bg_shape_5.fill.fore_color.rgb = DARK_BG
    bg_shape_5.line.fill.background()
    
    # Title & Subtitle box
    title_box_5 = slide_5.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.5))
    tf_5 = title_box_5.text_frame
    tf_5.word_wrap = True
    
    p_badge_5 = tf_5.paragraphs[0]
    p_badge_5.text = "SLIDE 5  |  TECHNOLOGY STACK"
    p_badge_5.font.size = Pt(12)
    p_badge_5.font.bold = True
    p_badge_5.font.color.rgb = PURPLE
    p_badge_5.font.name = "Arial"
    p_badge_5.space_after = Pt(10)
    
    p_title_5 = tf_5.add_paragraph()
    p_title_5.text = "Technologies Used"
    p_title_5.font.size = Pt(36)
    p_title_5.font.bold = True
    p_title_5.font.color.rgb = WHITE
    p_title_5.font.name = "Arial"
    
    # Question text box
    q_box = slide_5.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(0.6))
    tf_q = q_box.text_frame
    tf_q.word_wrap = True
    p_q = tf_q.paragraphs[0]
    p_q.text = "What technologies, frameworks, and tools were used and why were they selected for this solution?"
    p_q.font.size = Pt(16)
    p_q.font.italic = True
    p_q.font.color.rgb = LIGHT_GRAY
    p_q.font.name = "Arial"
    
    # Left Column: Technologies List 1
    left_col_5 = slide_5.shapes.add_textbox(Inches(1.0), Inches(3.0), Inches(5.3), Inches(4.0))
    tf_left_5 = left_col_5.text_frame
    tf_left_5.word_wrap = True
    
    tech_bullets_left = [
        ("FAISS (Dense Search)", "Selected for high-performance, fast vector similarity search on dense section embeddings, scaling efficiently to millions of profiles."),
        ("BM25 (Sparse Search)", "Selected for sparse exact keyword matching, ensuring specific developer skills and ontology terms are matched precisely without semantic dilution."),
        ("XGBoost (LTR Student)", "Selected as a lightweight, super-fast LTR student model that handles complex non-linear combinations of 46 structural features in milliseconds.")
    ]
    for title, desc in tech_bullets_left:
        p_t = tf_left_5.add_paragraph() if tf_left_5.paragraphs[0].text else tf_left_5.paragraphs[0]
        p_t.text = "• " + title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = PURPLE
        p_t.font.name = "Arial"
        p_t.space_after = Pt(2)
        
        p_d = tf_left_5.add_paragraph()
        p_d.text = "  " + desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = LIGHT_GRAY
        p_d.font.name = "Arial"
        p_d.space_after = Pt(10)
        
    # Right Column: Technologies List 2
    right_col_5 = slide_5.shapes.add_textbox(Inches(7.0), Inches(3.0), Inches(5.3), Inches(4.0))
    tf_right_5 = right_col_5.text_frame
    tf_right_5.word_wrap = True
    
    tech_bullets_right = [
        ("SentenceTransformers (Embeddings)", "BGE-base-en-v1.5 selected for dense semantic bi-encoder embeddings; ms-marco-MiniLM-L-6-v2 selected for high-fidelity cross-encoder query-resume deep text matching."),
        ("SQLite (Feature DB)", "Selected as a lightweight, serverless relational database to store candidate metadata and retrieve computed candidate features instantaneously."),
        ("Streamlit (UI Portal)", "Selected for building a clean, responsive, and lightweight web interface for recruiters to interactively upload job descriptions and explore rankings.")
    ]
    for title, desc in tech_bullets_right:
        p_t = tf_right_5.add_paragraph() if tf_right_5.paragraphs[0].text else tf_right_5.paragraphs[0]
        p_t.text = "• " + title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = BLUE
        p_t.font.name = "Arial"
        p_t.space_after = Pt(2)
        
        p_d = tf_right_5.add_paragraph()
        p_d.text = "  " + desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = LIGHT_GRAY
        p_d.font.name = "Arial"
        p_d.space_after = Pt(10)
        
    # Save the presentation
    output_path = "Candidate_Ranking_System_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
