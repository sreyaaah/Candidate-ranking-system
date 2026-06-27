import sys
import os
import streamlit as st
import pandas as pd
import sqlite3
from docx import Document

sys.path.append("src")
from search_hybrid import get_hybrid_scores

st.set_page_config(layout="wide", page_title="INDIA.RUNS Candidate Ranking")
st.title("AI Candidate Ranking System")

required_skills = [
    "python",
    "react",
    "node.js",
    "docker",
    "sql"
]

@st.cache_data
def get_rankings():
    # Load JD
    doc = Document("data/jobs/job_description.docx")
    jd_text = "\n".join([para.text for para in doc.paragraphs])
    
    # Get Hybrid RRF scores
    top_n = 2000
    hybrid_results = get_hybrid_scores(jd_text, top_n=top_n)
    
    candidate_ids = [res[0] for res in hybrid_results]
    
    # Fetch DB features
    conn = sqlite3.connect("data/candidates.db")
    cursor = conn.cursor()
    placeholders = ",".join(["?"] * len(candidate_ids))
    
    query = f"""
        SELECT 
            c.id, c.skills, 
            f.yoe, f.github_score, f.education_tier, f.notice_period_days, f.ai_years
        FROM candidates c
        LEFT JOIN candidate_features f ON c.id = f.candidate_id
        WHERE c.id IN ({placeholders})
    """
    cursor.execute(query, candidate_ids)
    
    features_map = {}
    for row in cursor.fetchall():
        features_map[row[0]] = {
            "skills": row[1].lower() if row[1] else "",
            "yoe": row[2] or 0.0,
            "github_score": row[3] or 0.0,
            "education_tier": row[4] or 3,
            "notice_period": row[5] or 0,
            "ai_years": row[6] or 0.0
        }
    conn.close()
    
    results = []
    
    for cand_id, rrf_score in hybrid_results:
        feats = features_map.get(cand_id, {})
        text = feats.get("skills", "")
        
        # Calculate Skill Score
        matched = sum(1 for skill in required_skills if skill in text)
        skill_score = matched / len(required_skills)
        
        # Calculate Final Score
        scaled_rrf = rrf_score * 30
        final_score = (0.7 * scaled_rrf) + (0.3 * skill_score)
        
        # Apply Stage 3 Heuristics
        yoe = feats.get("yoe", 0.0)
        github = feats.get("github_score", 0.0)
        tier = feats.get("education_tier", 3)
        notice = feats.get("notice_period", 0)
        
        final_score += min(yoe, 10) * 0.01 
        final_score += (github / 100.0) * 0.1
        if tier == 1: final_score += 0.05
        if notice > 60: final_score -= 0.10
            
        results.append({
            "Candidate ID": cand_id,
            "Heuristic Score": round(final_score, 4),
            "Hybrid RRF": round(scaled_rrf, 4),
            "Skill Score": round(skill_score, 2),
            "Total YoE": round(yoe, 1),
            "GitHub Score": round(github, 1),
            "Edu Tier": tier,
            "Notice Period": notice
        })
        
    df = pd.DataFrame(results)
    df = df.sort_values(by="Heuristic Score", ascending=False).reset_index(drop=True)
    return df

st.markdown("### Job Description Features")
st.write(f"**Required Skills:** {', '.join(required_skills)}")

with st.spinner('Calculating hybrid retrieval and engineering features...'):
    df = get_rankings()

st.subheader(f"Top {len(df)} Ranked Candidates")
st.dataframe(df, use_container_width=True)